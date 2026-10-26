"""Genera el entregable FINAL de ACA 1 (Proyecto I) en Word, aplicando los hallazgos de
critica-docente.md sobre el borrador de ../trabajo/. Igual que el borrador, siguiendo la plantilla oficial
CUN destilada en ../../materiales/plantilla-anteproyecto.txt (formato: carta, margenes
3 cm, Times New Roman 12, interlineado 1.5) y el alcance de la rubrica en
../planificacion/criterios-aceptacion.md (hasta "Alcances y Limitaciones" -- sin marco
referencial, metodologia ni cronograma, que son de ACA 2/ACA 3).
Fuente de contenido: ../planificacion/solucion.md
Salida: ACA1_Anteproyecto_Apollo_Services.docx (y su conversion a PDF via soffice)

Tabla de contenido estatica: las paginas salen de toc_paginas.json, que escribe
calcular_toc.py tras una primera pasada (generar -> PDF -> ubicar titulos -> regenerar),
para que el PDF muestre los numeros sin depender de actualizar campos en Word.
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import json
import os

OUT_DIR = os.path.dirname(os.path.abspath(__file__))
OUT_PATH = os.path.join(OUT_DIR, "ACA1_Anteproyecto_Apollo_Services.docx")
TOC_PATH = os.path.join(OUT_DIR, "toc_paginas.json")
TOC_PAGES = json.load(open(TOC_PATH, encoding="utf-8")) if os.path.exists(TOC_PATH) else {}

INTEGRANTES = [
    "Jonathan Adonay Rubio Jimenez",
    "Jhoan Jesus Cardozo Gelvis",
]

doc = Document()

section = doc.sections[0]
section.top_margin = Cm(3.0)
section.bottom_margin = Cm(3.0)
section.left_margin = Cm(3.0)
section.right_margin = Cm(3.0)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.first_line_indent = Cm(1.27)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.space_after = Pt(0)

h1 = doc.styles["Heading 1"]
h1.font.name = "Times New Roman"
h1.font.size = Pt(12)
h1.font.bold = True
h1.font.color.rgb = RGBColor(0, 0, 0)
h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
h1.paragraph_format.space_before = Pt(12)
h1.paragraph_format.space_after = Pt(6)
h1.paragraph_format.first_line_indent = Cm(0)

h2 = doc.styles["Heading 2"]
h2.font.name = "Times New Roman"
h2.font.size = Pt(12)
h2.font.bold = True
h2.font.color.rgb = RGBColor(0, 0, 0)
h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
h2.paragraph_format.space_before = Pt(10)
h2.paragraph_format.space_after = Pt(4)
h2.paragraph_format.first_line_indent = Cm(0)


def fix_style_fonts(style):
    """Los estilos de título traen fuentes de tema (Calibri) que ganan sobre
    font.name; se eliminan para que quede Times New Roman en todos los scripts."""
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in list(rfonts.attrib):
        if attr.endswith("Theme"):
            del rfonts.attrib[attr]
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), "Times New Roman")


for st in (normal, h1, h2):
    fix_style_fonts(st)


def heading(text, level, new_page=False):
    """Título; new_page usa page_break_before en vez de un párrafo de salto,
    que dejaba una página en blanco cuando la anterior quedaba llena."""
    hd = doc.add_heading(text, level=level)
    hd.paragraph_format.page_break_before = new_page
    return hd


def add_page_number_field(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)


header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hp.paragraph_format.first_line_indent = Cm(0)
add_page_number_field(hp)


def p(text="", bold=False, italic=False, align=None, indent0=False):
    para = doc.add_paragraph()
    if indent0:
        para.paragraph_format.first_line_indent = Cm(0)
    if align is not None:
        para.paragraph_format.alignment = align
    if text:
        run = para.add_run(text)
        run.bold = bold
        run.italic = italic
    return para


def add_reference(text):
    """Referencia APA 7 con sangria francesa (hanging indent). Los tramos entre
    asteriscos (*Revista, vol*) van en cursiva, como exige APA 7."""
    para = doc.add_paragraph()
    para.paragraph_format.first_line_indent = Cm(-1.27)
    para.paragraph_format.left_indent = Cm(1.27)
    para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_after = Pt(8)
    for i, chunk in enumerate(text.split("*")):
        if chunk:
            para.add_run(chunk).italic = i % 2 == 1
    return para


# ---------------------------------------------------------------------------
# PORTADA
# ---------------------------------------------------------------------------
for _ in range(3):
    p()
p("Especialización en Inteligencia Artificial", align=WD_ALIGN_PARAGRAPH.CENTER, indent0=True)
p("Línea de profundización: Uso y adaptación de IA para entornos productivos",
  align=WD_ALIGN_PARAGRAPH.CENTER, indent0=True)
p()
p("Título del Proyecto:", align=WD_ALIGN_PARAGRAPH.CENTER, indent0=True)
p("Efecto de la clasificación conversacional con inteligencia artificial en la "
  "conversión de prospectos a citas en una pyme bogotana de servicios",
  bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent0=True)
p()
p("Presentado por:", align=WD_ALIGN_PARAGRAPH.CENTER, indent0=True)
for nombre in INTEGRANTES:
    p(nombre, align=WD_ALIGN_PARAGRAPH.CENTER, indent0=True)
p()
p("Bajo la dirección de:", align=WD_ALIGN_PARAGRAPH.CENTER, indent0=True)
p("María Fernanda Rivera Sanclemente (Formador Virtual, Proyecto I)",
  align=WD_ALIGN_PARAGRAPH.CENTER, indent0=True)
p()
p("Bogotá, D. C.", align=WD_ALIGN_PARAGRAPH.CENTER, indent0=True)

# ---------------------------------------------------------------------------
# TABLA DE CONTENIDO (estatica, ver docstring)
# ---------------------------------------------------------------------------
TOC_ENTRIES = [
    ("Resumen", 1), ("Abstract", 1),
    ("Declaración de originalidad y autonomía", 2),
    ("Declaración de exoneración de responsabilidad", 2),
    ("Introducción", 1), ("Planteamiento del problema", 1),
    ("Diagnóstico de la situación actual", 2),
    ("Predicción o pronóstico de la situación descrita", 2),
    ("Formulación o Pregunta de Investigación", 2),
    ("Objetivos", 1), ("Objetivo general", 2), ("Objetivos específicos", 2),
    ("Justificación", 1), ("Alcances y Limitaciones del Proyecto", 1),
    ("Referencias bibliográficas", 1),
]
heading("Tabla de contenido", 1, new_page=True)
TEXT_WIDTH = Cm(21.59 - 6.0)
for titulo, nivel in TOC_ENTRIES:
    tp = p(indent0=True)
    tp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    tp.paragraph_format.left_indent = Cm(0 if nivel == 1 else 1.0)
    tp.paragraph_format.tab_stops.add_tab_stop(TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT,
                                               WD_TAB_LEADER.DOTS)
    tp.add_run(f"{titulo}\t{TOC_PAGES.get(titulo, '00')}")

# ---------------------------------------------------------------------------
# RESUMEN / ABSTRACT
# ---------------------------------------------------------------------------
heading("Resumen", 1, new_page=True)
p(
    "Este anteproyecto formula y delimita el problema de investigación de un "
    "trabajo de grado de especialización vinculado al ecosistema de inteligencia "
    "artificial Apollo Services, un CRM comercial multicanal en fase de diseño. El "
    "problema identificado es la pérdida de conversión comercial que enfrentan las "
    "pymes de servicios de Bogotá al atender múltiples canales digitales (WhatsApp, "
    "Instagram, Facebook) sin un sistema de clasificación automatizado de "
    "contactos, lo que retrasa las respuestas y hace depender el agendamiento de "
    "citas de la confirmación manual. A partir de este diagnóstico se formula la "
    "pregunta de investigación: ¿en qué medida la implementación de un sistema de "
    "clasificación conversacional basado en inteligencia artificial afecta la tasa "
    "de conversión de prospectos a citas comerciales confirmadas en una pyme de "
    "servicios de Bogotá? El objetivo general es evaluar ese efecto mediante un "
    "estudio de caso, apoyado en cuatro objetivos específicos que van desde "
    "caracterizar el proceso comercial actual hasta comparar la conversión antes y "
    "después de la intervención. La justificación se sustenta en evidencia "
    "académica reciente sobre adopción de inteligencia artificial en pequeñas y "
    "medianas empresas y sobre el papel de los chatbots conversacionales en la "
    "conversión comercial. Este documento corresponde a la primera entrega "
    "acumulativa (ACA 1) de la asignatura: delimita el problema y las piezas "
    "iniciales del anteproyecto, sin abordar todavía el marco referencial ni el "
    "diseño metodológico, que se desarrollarán en las siguientes entregas."
)
p("Palabras clave: inteligencia artificial, clasificación conversacional, "
  "conversión comercial, pymes de servicios, anteproyecto de investigación.",
  italic=False)

heading("Abstract", 1)
p(
    "This pre-proposal formulates and delimits the research problem of a "
    "specialization capstone project linked to Apollo Services, a "
    "multichannel AI-powered commercial CRM ecosystem at the design stage. The "
    "identified problem is the loss of commercial conversion faced by service "
    "SMEs in Bogotá when handling multiple digital channels (WhatsApp, Instagram, "
    "Facebook) without an automated contact-classification system, which delays "
    "responses and makes appointment scheduling dependent on manual confirmation. "
    "Based on this diagnosis, the research question is posed: to what extent does "
    "the implementation of an AI-based conversational classification system "
    "affect the conversion rate of prospects into confirmed commercial "
    "appointments in a service SME in Bogotá? The general objective is to evaluate "
    "this effect through a case study, supported by four specific objectives "
    "ranging from characterizing the current commercial process to comparing "
    "conversion before and after the intervention. The "
    "justification is grounded in recent academic evidence on artificial "
    "intelligence adoption in small and medium enterprises and on the role of "
    "conversational chatbots in commercial conversion. This document corresponds "
    "to the first cumulative deliverable (ACA 1) of the course: it delimits the "
    "problem and the initial pieces of the pre-proposal, without yet addressing "
    "the referential framework or the methodological design, which will be "
    "developed in the following deliverables."
)
p("Keywords: artificial intelligence, conversational classification, commercial "
  "conversion, service SMEs, research proposal.")

# ---------------------------------------------------------------------------
# PRELIMINARES — DECLARACIONES
# ---------------------------------------------------------------------------
heading("Declaración de originalidad y autonomía", 2, new_page=True)
p(
    "Declaramos bajo la gravedad del juramento, que hemos escrito el presente "
    "proyecto, en la propuesta de solución a una problemática en el campo de "
    "conocimientos del programa de Especialización en Inteligencia Artificial, por "
    "nuestra propia cuenta y que, por lo tanto, su contenido es original."
)
p(
    "Declaramos que hemos indicado clara y precisamente todas las fuentes directas "
    "e indirectas de información y que no ha sido entregado a ninguna otra "
    "institución con fines de calificación o publicación."
)
for nombre in INTEGRANTES:
    p()
    p("_______________________________", align=WD_ALIGN_PARAGRAPH.CENTER, indent0=True)
    p(nombre, align=WD_ALIGN_PARAGRAPH.CENTER, indent0=True)
p()
p("Firmado en Bogotá, D. C.", align=WD_ALIGN_PARAGRAPH.CENTER, indent0=True)

heading("Declaración de exoneración de responsabilidad", 2)
p(
    "Declaramos que la responsabilidad intelectual del presente trabajo es "
    "exclusivamente de sus autores. La Corporación Unificada Nacional de Educación "
    "Superior – CUN no se hace responsable de contenidos, opiniones o ideologías "
    "expresadas total o parcialmente en él."
)
for nombre in INTEGRANTES:
    p()
    p("_______________________________", align=WD_ALIGN_PARAGRAPH.CENTER, indent0=True)
    p(nombre, align=WD_ALIGN_PARAGRAPH.CENTER, indent0=True)
p()
p("Firmado en Bogotá, D. C.", align=WD_ALIGN_PARAGRAPH.CENTER, indent0=True)

# ---------------------------------------------------------------------------
# INTRODUCCIÓN
# ---------------------------------------------------------------------------
heading("Introducción", 1, new_page=True)
p(
    "Las organizaciones de servicios que operan con equipos comerciales pequeños "
    "enfrentan un reto creciente: atender simultáneamente varios canales digitales "
    "de contacto (WhatsApp, Instagram, Facebook) sin perder oportunidades de venta "
    "por falta de tecnificación. La adopción de inteligencia artificial (IA) en "
    "pequeñas y medianas empresas (pymes) ha sido documentada como un proceso "
    "heterogéneo, condicionado por la madurez digital, la capacidad "
    "organizacional y la disponibilidad de recursos de cada empresa (Popa et al., "
    "2025; Sánchez et al., 2025). A pesar de estas barreras, la literatura "
    "reciente muestra que la IA aplicada a la atención al cliente —particularmente "
    "los chatbots conversacionales— puede mejorar tanto la experiencia del "
    "consumidor como las tasas de conversión, al responder con mayor rapidez y "
    "consistencia que un equipo humano sobrecargado (Ghosh et al., 2024; Ramki et "
    "al., 2024)."
)
p(
    "En paralelo, estudios sobre adopción sostenible de IA en pymes advierten que "
    "el beneficio de estas tecnologías no es automático: depende de que la "
    "organización desarrolle capacidades dinámicas que le permitan integrar la "
    "herramienta a su proceso comercial real, y no solo adquirirla como una "
    "solución aislada (Cimino et al., 2025). Este anteproyecto retoma esa "
    "discusión aplicándola a un caso concreto: Apollo Services, un ecosistema de "
    "agentes de IA conversacional diseñado para unificar la gestión comercial "
    "multicanal de pymes de servicios, documentado previamente como solución de "
    "negocio independiente de este trabajo académico. A partir de ese caso, se "
    "busca formular y delimitar el problema de investigación que permitirá, en "
    "fases posteriores de la especialización (Proyecto II), medir el efecto real "
    "de este tipo de intervención tecnológica sobre la conversión comercial."
)
p(
    "En concreto, el problema que aborda este anteproyecto es la pérdida de "
    "conversión comercial en pymes de servicios que atienden varios canales "
    "digitales sin un sistema automatizado de clasificación de contactos. De allí se "
    "desprende la pregunta que orienta el estudio: ¿en qué medida un sistema de "
    "clasificación conversacional basado en IA afecta la tasa de conversión de "
    "prospectos a citas comerciales confirmadas? En consecuencia, el objetivo general "
    "consiste en evaluar ese efecto en una pyme de servicios de Bogotá, y su "
    "justificación descansa en la necesidad práctica de las pymes y en el vacío de "
    "evidencia aplicada que señala la literatura. El documento presenta, en ese "
    "orden, el planteamiento del problema, los objetivos, la justificación y los "
    "alcances y limitaciones del proyecto."
)

# ---------------------------------------------------------------------------
# PLANTEAMIENTO DEL PROBLEMA
# ---------------------------------------------------------------------------
heading("Planteamiento del problema", 1)

heading("Diagnóstico de la situación actual", 2)
p(
    "En las pymes de servicios, los equipos comerciales pequeños (entre 2 y 5 "
    "personas, el segmento al que se dirige Apollo Services) atienden manualmente "
    "conversaciones provenientes de múltiples redes sociales. Cada canal se "
    "revisa por separado, sin una bandeja unificada ni un criterio homogéneo para "
    "distinguir un contacto meramente curioso (Seguidor) de uno con intención real "
    "de compra (Prospecto). Esta atención dispersa genera tres consecuencias "
    "concretas: (a) contactos que no reciben respuesta a tiempo, (b) ausencia de "
    "un punto único de gestión y trazabilidad de las conversaciones, y (c) pérdida "
    "del cierre comercial porque el agendamiento de citas queda sujeto a que "
    "alguien confirme manualmente en un calendario compartido."
)
p(
    "El problema de investigación no es, por tanto, “cómo construir una "
    "plataforma de inteligencia artificial” —la plataforma es la posible "
    "solución—, sino la pérdida de conversión comercial por ausencia de un "
    "sistema de clasificación y seguimiento automatizado de contactos "
    "multicanal. Esta descripción es consistente con lo que la literatura reciente "
    "identifica como barrera estructural de las pymes para beneficiarse de la IA: "
    "no es la tecnología en sí lo escaso, sino la capacidad organizacional para "
    "integrarla a procesos ya existentes (Cimino et al., 2025; Sánchez et al., "
    "2025)."
)
p(
    "El problema se observa a través de tres indicadores, que el estudio "
    "medirá en el caso seleccionado: (a) el tiempo de primera respuesta a un "
    "contacto nuevo, (b) la proporción de contactos que no reciben respuesta y (c) "
    "la tasa de conversión de prospecto a cita confirmada, que es la variable "
    "central de la investigación. Los dos primeros explican la pérdida; el "
    "tercero la cuantifica."
)
p(
    "Delimitación: el estudio se circunscribe a una pyme del sector servicios "
    "ubicada en Bogotá, con un equipo comercial de 2 a 5 personas que atiende "
    "WhatsApp, Instagram y Facebook, durante el periodo de ejecución de Proyecto II."
)

heading("Predicción o pronóstico de la situación descrita", 2)
p(
    "Si esta situación no se interviene, es previsible que la pérdida de "
    "prospectos se mantenga o se agrave a medida que el volumen de contactos "
    "crezca: cada canal adicional que la organización sume sin un sistema de "
    "clasificación centralizado multiplica el riesgo de que un contacto con "
    "intención real de compra quede sin respuesta a tiempo. A mediano plazo, esto "
    "se traduce en una desventaja competitiva frente a organizaciones similares "
    "que sí automaticen la calificación y el seguimiento de sus contactos, y en un "
    "desgaste del equipo comercial, que debe sostener manualmente una carga de "
    "atención que crece más rápido que su capacidad de respuesta."
)

heading("Formulación o Pregunta de Investigación", 2)
p(
    "¿En qué medida la implementación de un sistema de clasificación "
    "conversacional basado en inteligencia artificial, que unifica los canales de "
    "contacto y automatiza la calificación de Seguidor/Prospecto, afecta la tasa "
    "de conversión de prospectos a citas comerciales confirmadas en una pyme de "
    "servicios de Bogotá?"
)
p(
    "La pregunta se centra en una sola variable dependiente, la tasa de "
    "conversión de prospecto a cita confirmada, observada en dos momentos: sin "
    "el sistema de clasificación y con él. La conversión se calcula como la "
    "proporción de prospectos que llegan a una cita confirmada sobre el total de "
    "prospectos clasificados en cada periodo, lo que la hace observable con los "
    "registros de la propia organización."
)

# ---------------------------------------------------------------------------
# OBJETIVOS
# ---------------------------------------------------------------------------
heading("Objetivos", 1)

heading("Objetivo general", 2)
p(
    "Evaluar el efecto de la implementación de un sistema de clasificación "
    "conversacional basado en inteligencia artificial en la tasa de conversión de "
    "prospectos a citas comerciales confirmadas, en una pyme de servicios de Bogotá."
)

heading("Objetivos específicos", 2)
for n, texto in enumerate([
    "Caracterizar el proceso comercial actual de la pyme de servicios, sin "
    "sistema de clasificación con IA, con la identificación de los puntos donde "
    "se pierden prospectos.",
    "Definir las condiciones de implementación del sistema de clasificación "
    "conversacional (bandeja unificada, clasificación Seguidor/Prospecto y "
    "agendamiento de Apollo Services) que constituyen la intervención.",
    "Medir la tasa de conversión de prospectos a citas confirmadas antes y "
    "después de la intervención.",
    "Comparar las tasas de conversión obtenidas en ambos periodos para "
    "determinar la magnitud del efecto y las condiciones organizacionales que lo "
    "favorecen o lo limitan.",
], start=1):
    bp = p(indent0=True)
    bp.paragraph_format.left_indent = Cm(1.27)
    bp.paragraph_format.first_line_indent = Cm(-0.63)
    bp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    bp.add_run(f"{n}. {texto}")

# ---------------------------------------------------------------------------
# JUSTIFICACIÓN
# ---------------------------------------------------------------------------
heading("Justificación", 1)
p(
    "La pertinencia de este anteproyecto se sustenta en tres planos. Primero, en "
    "el plano práctico: la pérdida de prospectos por atención multicanal "
    "desorganizada es un problema de negocio real y vigente, no hipotético. "
    "Segundo, en el plano académico: la literatura reciente coincide en que la "
    "adopción de IA en pymes depende de variables organizacionales más que "
    "puramente tecnológicas (Popa et al., 2025; Sánchez et al., 2025); sin "
    "embargo, la literatura revisada se concentra en los determinantes de la "
    "adopción y no en su efecto directo sobre la conversión comercial, y ninguno de "
    "los estudios consultados se ubica en pymes de servicios latinoamericanas. "
    "Este anteproyecto aporta un caso delimitado en ese contexto. Tercero, en el plano profesional: estudios "
    "sobre chatbots conversacionales en atención al cliente muestran mejoras "
    "medibles en satisfacción y conversión cuando la herramienta está bien "
    "integrada al proceso de negocio (Ghosh et al., 2024; Ramki et al., 2024), lo "
    "que valida que medir ese efecto en un caso real tiene valor replicable para "
    "otras pymes con restricciones similares de presupuesto tecnológico."
)
p(
    "Desde el punto de vista metodológico, el estudio aporta un procedimiento "
    "replicable para medir la conversión comercial antes y después de introducir "
    "un sistema de IA, apoyado en los registros que la propia organización ya "
    "genera. Los beneficiarios directos son los equipos comerciales de pymes de "
    "servicios, que dispondrían de evidencia para decidir si automatizar la "
    "clasificación de sus contactos. De forma indirecta, el trabajo alimenta la "
    "línea de profundización en uso y adaptación de IA para entornos productivos "
    "con un caso aplicado y medible."
)

# ---------------------------------------------------------------------------
# ALCANCES Y LIMITACIONES DEL PROYECTO
# ---------------------------------------------------------------------------
heading("Alcances y Limitaciones del Proyecto", 1)
p(
    "Alcances: el proyecto mide el efecto del sistema de clasificación "
    "conversacional sobre una sola variable, la tasa de conversión de prospectos a "
    "citas confirmadas, en una pyme de servicios de Bogotá, comparando un periodo "
    "sin el sistema y otro con él. Se compromete a entregar la caracterización del "
    "proceso comercial, las condiciones de la intervención, la medición en ambos "
    "periodos y el análisis comparativo. No evalúa ingresos, satisfacción del "
    "cliente ni retorno de la inversión. Este documento cubre únicamente la "
    "formulación del problema; el marco referencial, la metodología y el "
    "cronograma se desarrollan en las entregas siguientes, y la ejecución en "
    "Proyecto II."
)
p(
    "Limitaciones: primero, el estudio se centra en un único caso de "
    "intervención (Apollo Services aplicado a una pyme de servicios), por lo que "
    "sus hallazgos no son generalizables sin validación adicional en otros "
    "sectores; segundo, la medición de la variable de conversión depende de que "
    "la organización lleve un registro consistente de sus interacciones antes y "
    "después de la intervención, requisito operativo que debe confirmarse con la "
    "empresa caso antes de avanzar a la fase de desarrollo metodológico; tercero, "
    "la propuesta se limita a los canales ya cubiertos por el sistema (WhatsApp, "
    "Instagram y Facebook) y no evalúa canales adicionales fuera del alcance "
    "documentado del proyecto."
)
p(
    "Viabilidad: el proyecto es académicamente viable porque el caso de negocio "
    "ya está documentado a nivel de requerimientos y diseño técnico, y porque se "
    "apoya en literatura reciente y accesible. Es operativamente viable porque los "
    "tres módulos que requiere la intervención (bandeja unificada, clasificación y "
    "agendamiento) tienen su diseño cerrado, de modo que la ejecución no parte de "
    "cero; su condición es contar con una pyme dispuesta a facilitar sus registros "
    "comerciales. Es temporalmente viable porque se limita a una sola variable "
    "medible, en línea con la recomendación de trabajar con una o dos variables "
    "dado el tiempo de 12 semanas de cada curso."
)

# ---------------------------------------------------------------------------
# REFERENCIAS BIBLIOGRÁFICAS
# ---------------------------------------------------------------------------
heading("Referencias bibliográficas", 1, new_page=True)

add_reference(
    "Cimino, A., Corvello, V., Troise, C., Thomas, A., & Tani, M. (2025). "
    "Artificial intelligence adoption for sustainable growth in SMEs: An extended "
    "dynamic capability framework. *Corporate Social Responsibility and "
    "Environmental Management, 32*(5), 6120–6138. https://doi.org/10.1002/csr.70019"
)
add_reference(
    "Ghosh, S., Ness, S., & Salunkhe, S. (2024). The role of AI enabled chatbots "
    "in omnichannel customer service. *Journal of Engineering Research and "
    "Reports, 26*(6), 327–345. https://doi.org/10.9734/jerr/2024/v26i61184"
)
add_reference(
    "Popa, R.-G., Popa, I.-C., Ciocodeică, D.-F., & Mihălcescu, H. (2025). "
    "Modeling AI adoption in SMEs for sustainable innovation: A PLS-SEM approach "
    "integrating TAM, UTAUT2, and contextual drivers. *Sustainability, 17*(15), "
    "6901. https://doi.org/10.3390/su17156901"
)
add_reference(
    "Ramki, R., Gopi, V., Vidya, R., Markan, R., Natarajan, S., & Rajalakshmi, M. "
    "(2024). AI-powered chatbots in customer service: Impact on brand loyalty "
    "and conversion rates. *Economic Sciences, 20*(2), 190–203. "
    "https://doi.org/10.69889/vs5gtv52"
)
add_reference(
    "Sánchez, E., Calderón, R., & Herrera, F. (2025). Artificial intelligence "
    "adoption in SMEs: Survey based on TOE–DOI framework, primary methodology "
    "and challenges. *Applied Sciences, 15*(12), 6465. "
    "https://doi.org/10.3390/app15126465"
)

doc.save(OUT_PATH)
print(f"Documento generado: {OUT_PATH}")
